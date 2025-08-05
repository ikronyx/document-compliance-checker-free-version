# Handles file parsing
import pdfplumber
import docx
import openpyxl
import extract_msg
import eml_parser
import os

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_pdf(file_path)
    elif ext == ".docx":
        return extract_docx(file_path)
    elif ext == ".xlsx":
        return extract_excel(file_path)
    elif ext == ".msg":
        return extract_msg_file(file_path)
    elif ext == ".eml":
        return extract_eml_file(file_path)
    return ""

def extract_pdf(path):
    text = ""
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text += page.extract_text() + "\n"
    return text

def extract_docx(path):
    doc = docx.Document(path)
    return "\n".join([para.text for para in doc.paragraphs])

def extract_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    text = ""
    for sheet in wb:
        for row in sheet.iter_rows(values_only=True):
            line = " ".join([str(cell) if cell is not None else "" for cell in row])
            text += line + "\n"
    return text

def extract_msg_file(path):
    msg = extract_msg.Message(path)
    return msg.body or ""

def extract_eml_file(path):
    with open(path, "rb") as f:
        eml_data = eml_parser.eml_parser.decode_email_b(f.read(), include_raw_body=True)
    return eml_data.get("body", {}).get("plain", [""])[0]
